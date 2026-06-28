import docx

doc = docx.Document(r'C:\Users\Atish Gadade\OneDrive\Desktop\pdf to json\District Name.docx')
print("Paragraphs:")
for p in doc.paragraphs:
    if p.text.strip():
        print(p.text.strip())

print("\nTables:")
for table in doc.tables:
    for row in table.rows:
        row_data = [cell.text.strip() for cell in row.cells]
        print(" | ".join(row_data))
